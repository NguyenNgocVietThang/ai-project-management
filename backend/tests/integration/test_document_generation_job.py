"""Integration test D2: job sinh tài liệu SRS/BRD theo nhóm mục với provider và storage giả."""
import io
import json

import pytest
from docx import Document as DocxDocument
from fastapi import HTTPException
from sqlalchemy import func, insert, select
from sqlalchemy.ext.asyncio import async_sessionmaker

from app.core.job_runner import JobRunner
from app.models.associations import project_members
from app.models.change_request import ChangeRequest, CRStatus
from app.models.generated_document import GeneratedDocument
from app.models.job import Job
from app.models.project import Project
from app.models.requirement_brief import RequirementBrief
from app.models.task import Task
from app.models.user import User
from app.services import document_service
from app.services.ai.document_writer import plan_section_groups
from app.services.ai.job_context import active_ai_context, reserve_ai_call
from app.services.ai.parsing import USER_INPUT_FENCE
from app.services.document_service import request_document
from app.services.documents.templates import get_template
from app.services.durable_handlers import HANDLERS
from app.services.storage_service import StorageError
from tests.integration.test_requirement_clarify_job import exhaust_quota

ITEMS = {
    "goals": {"status": "clear", "value": "GIA_TRI_CLEAR bán cà phê online", "source": "user", "proposed_assumption": None},
    "stakeholders": {"status": "clear", "value": "Chủ quán", "source": "ai", "proposed_assumption": None},
    "in_scope": {"status": "assumed", "value": "GIA_TRI_ASSUMED web và mobile", "source": "ai_assumption_confirmed", "proposed_assumption": None},
    "out_of_scope": {"status": "skipped", "value": None, "source": None, "proposed_assumption": None},
    "key_features": {"status": "clear", "value": "Đặt hàng", "source": "user", "proposed_assumption": None},
    "constraints": {"status": "skipped", "value": None, "source": None, "proposed_assumption": None},
    "success_criteria": {"status": "missing", "value": None, "source": None, "proposed_assumption": "DE_XUAT_BIA_DAT"},
}


def sample_value(section):
    if section.kind == "table":
        return [{column.key: f"{section.key}-{column.key}" for column in section.columns}]
    if section.kind == "bullets":
        return [f"{section.key} một", f"{section.key} hai"]
    return [f"Đoạn văn {section.key}."]


class FakeProvider:
    """Provider giả theo giao thức checkpoint + quota; trả nội dung hợp lệ cho nhóm được hỏi,
    kèm một mục ngoài mẫu và một mục của mẫu nhưng ngoài nhóm (phải bị loại)."""

    def __init__(self, template):
        self.template = template
        self.calls = []
        self.wrong_type = False

    async def generate_json(self, prompt, system_prompt="", task=None):
        context = active_ai_context.get()
        if context.job.staged_result:
            return json.loads(context.job.staged_result["text"])
        await context.mark_dispatched(before_dispatch=reserve_ai_call)
        body = json.loads(prompt.split(USER_INPUT_FENCE)[1].strip())
        group = body["sections_to_write"]
        self.calls.append((group, body, system_prompt))
        sections = {key: sample_value(self.template.section(key)) for key in group}
        sections["hacked_section"] = ["không thuộc mẫu"]
        outsider = next(key for key in self.template.content_section_keys if key not in group)
        sections[outsider] = sample_value(self.template.section(outsider))
        if self.wrong_type:
            sections[group[0]] = {"not": "valid"}
        payload = {"sections": sections}
        await context.stage_result({"text": json.dumps(payload)})
        return payload


class FakeStorage:
    def __init__(self, fail=False):
        self.objects = {}
        self.fail = fail

    async def put_bytes(self, key, data, content_type):
        if self.fail:
            raise StorageError("Supabase storage connection failed")
        self.objects[key] = (data, content_type)

    async def get_bytes(self, key):
        return self.objects[key]


def install(monkeypatch, template, storage):
    provider = FakeProvider(template)

    async def factory():
        return provider

    monkeypatch.setattr("app.services.ai.project_generator.get_ai_provider", factory)
    monkeypatch.setattr(document_service, "get_storage_service", lambda: storage)
    return provider


async def make_project(session, make_user, seed_roles, *, items=ITEMS, with_brief=True):
    users = {name: await make_user() for name in ("PM", "BA", "Member")}
    project = Project(name="Quán cà phê", pm_id=users["PM"].id)
    session.add(project)
    await session.flush()
    for name in users:
        await session.execute(insert(project_members).values(
            project_id=project.id, user_id=users[name].id, role_id=seed_roles[name].id,
        ))
    session.add(Task(project_id=project.id, name="TASK_THAT", is_critical=True))
    session.add(ChangeRequest(project_id=project.id, requested_by_id=users["BA"].id, title="CR_DA_DUYET", description="d", status=CRStatus.APPROVED))
    session.add(ChangeRequest(project_id=project.id, requested_by_id=users["BA"].id, title="CR_NHAP", description="d", status=CRStatus.DRAFT))
    brief = None
    if with_brief:
        brief = RequirementBrief(
            owner_id=users["PM"].id, project_id=project.id, idea="Ứng dụng bán cà phê",
            items=json.loads(json.dumps(items)), qa_history=[], status="consumed",
        )
        session.add(brief)
    await session.flush()
    return project, users, brief


async def drain(engine, limit=20):
    """Chạy lần lượt từng job tới khi hết; mỗi vòng là một "worker" mới (mô phỏng restart)."""
    factory = async_sessionmaker(engine, expire_on_commit=False)
    for _ in range(limit):
        runner = JobRunner(factory, HANDLERS)
        claimed = await runner.poll_once()
        await runner.stop()
        if claimed == 0:
            return


async def run_one(engine):
    runner = JobRunner(async_sessionmaker(engine, expire_on_commit=False), HANDLERS)
    claimed = await runner.poll_once()
    await runner.stop()
    return claimed


@pytest.mark.asyncio
async def test_generates_srs_in_groups_filters_foreign_sections_and_stores_private_docx(engine, session, make_user, seed_roles, monkeypatch):
    template = get_template("srs")
    storage = FakeStorage()
    provider = install(monkeypatch, template, storage)
    project, users, _ = await make_project(session, make_user, seed_roles)
    document, job = await request_document(session, project.id, users["BA"], "srs")
    assert job is not None and document.status == "queued"
    document_id, project_id = document.id, project.id
    await session.commit()
    await drain(engine)

    session.expire_all()
    document = await session.get(GeneratedDocument, document_id)
    groups = plan_section_groups(template)
    assert document.status == "ready", document.error_code
    assert len(provider.calls) == len(groups) >= 2
    assert [call[0] for call in provider.calls] == groups
    assert all(2 <= len(group) <= 4 or group is groups[-1] for group in groups)
    # Mục ngoài mẫu và ngoài nhóm bị loại, có cảnh báo; nội dung từng nhóm chỉ gồm mục của nhóm.
    for index, group in enumerate(groups):
        assert set(document.sections[str(index)]) == set(group)
    assert any("hacked_section" in warning for warning in document.warnings)
    assert any("không thuộc nhóm" in warning for warning in document.warnings)
    # Tệp lưu bucket private với key không đoán được.
    assert document.storage_key.startswith(f"documents/{project_id}/") and document.storage_key.endswith(".docx")
    assert len(document.storage_key.split("/")[-1]) == 32 + len(".docx")
    data, content_type = storage.objects[document.storage_key]
    assert document.file_size == len(data) and content_type.endswith("wordprocessingml.document")
    text = "\n".join(p.text for p in DocxDocument(io.BytesIO(data)).paragraphs)
    assert "hacked_section" not in text
    cells = "\n".join(cell.text for table in DocxDocument(io.BytesIO(data)).tables for row in table.rows for cell in row.cells)
    # Giả định/vấn đề mở sinh từ hồ sơ: in_scope là giả định, constraints/out_of_scope bỏ qua.
    assert "GIA_TRI_ASSUMED web và mobile" in cells and "Ràng buộc" in cells
    assert "DE_XUAT_BIA_DAT" not in cells

    # Prompt chỉ có dữ liệu hồ sơ clear/assumed (assumed tách riêng) và dữ liệu dự án thật.
    _, body, system_prompt = provider.calls[0]
    assert body["brief"]["requirements"]["goals"].startswith("GIA_TRI_CLEAR")
    assert body["brief"]["assumptions"] == {"in_scope": "GIA_TRI_ASSUMED web và mobile"}
    assert set(body["brief"]["unspecified_items"]) == {"out_of_scope", "constraints", "success_criteria"}
    assert "DE_XUAT_BIA_DAT" not in json.dumps(body, ensure_ascii=False)
    assert body["project"]["critical_path"] == ["TASK_THAT"]
    assert [cr["title"] for cr in body["project"]["approved_change_requests"]] == ["CR_DA_DUYET"]
    assert "Do NOT invent" in system_prompt and "Vietnamese" in system_prompt

    # Yêu cầu lại sau khi ready tạo phiên bản mới.
    await session.refresh(users["BA"])
    again, _ = await request_document(session, project_id, users["BA"], "SRS")
    assert again.version == 2 and again.id != document_id


@pytest.mark.asyncio
async def test_restart_midway_resumes_from_saved_groups_without_recalling(engine, session, make_user, seed_roles, monkeypatch):
    template = get_template("brd")
    storage = FakeStorage()
    provider = install(monkeypatch, template, storage)
    # BRD cần success_criteria: hồ sơ phải có mục này.
    items = json.loads(json.dumps(ITEMS))
    items["success_criteria"] = {"status": "clear", "value": "100 đơn/ngày", "source": "user", "proposed_assumption": None}
    project, users, _ = await make_project(session, make_user, seed_roles, items=items)
    document, job = await request_document(session, project.id, users["PM"], "brd", language="en")
    assert job is not None
    document_id = document.id
    await session.commit()
    groups = plan_section_groups(template)

    # Worker đầu chỉ kịp xong một nhóm rồi "chết"; job kế tiếp đã nằm trong hàng đợi.
    assert await run_one(engine) == 1
    session.expire_all()
    document = await session.get(GeneratedDocument, document_id)
    assert set(document.sections) == {"0"} and document.status == "generating"
    queued = await session.scalar(select(func.count()).select_from(Job).where(Job.status == "queued"))
    assert queued == 1 and len(provider.calls) == 1

    # Một job cũ bị chạy lại (trùng) không gọi provider và không ghi đè.
    first_job = await session.scalar(select(Job).where(Job.status == "succeeded"))
    first_job.status, first_job.lease_token = "queued", None
    await session.commit()
    await drain(engine)
    session.expire_all()
    document = await session.get(GeneratedDocument, document_id)
    assert document.status == "ready", document.error_code
    assert len(provider.calls) == len(groups)
    assert [call[0] for call in provider.calls] == groups
    assert "English" in provider.calls[0][2]


@pytest.mark.asyncio
async def test_provider_error_fails_then_retry_keeps_saved_groups(engine, session, make_user, seed_roles, monkeypatch):
    template = get_template("srs")
    storage = FakeStorage()
    provider = install(monkeypatch, template, storage)
    project, users, _ = await make_project(session, make_user, seed_roles)
    document, _ = await request_document(session, project.id, users["BA"], "srs")
    document_id, project_id, user_id = document.id, project.id, users["BA"].id
    await session.commit()
    assert await run_one(engine) == 1
    provider.wrong_type = True
    await drain(engine)
    session.expire_all()
    document = await session.get(GeneratedDocument, document_id)
    # Sai kiểu => lỗi provider, nhóm lỗi không được lưu.
    assert document.status == "failed" and document.error_code == "invalid_provider_response"
    assert set(document.sections) == {"0"}

    provider.wrong_type = False
    user = await session.get(User, user_id)
    retried, job = await request_document(session, project_id, user, "srs")
    assert retried.id == document_id and job is not None
    await session.commit()
    await drain(engine)
    session.expire_all()
    document = await session.get(GeneratedDocument, document_id)
    assert document.status == "ready"
    groups = plan_section_groups(template)
    # Nhóm 0 chỉ gọi một lần; nhóm 1 gọi hai lần (lần lỗi + lần chạy lại).
    called = [tuple(call[0]) for call in provider.calls]
    assert called.count(tuple(groups[0])) == 1
    assert len(provider.calls) == len(groups) + 1


@pytest.mark.asyncio
async def test_storage_error_fails_clearly_and_retry_does_not_call_ai(engine, session, make_user, seed_roles, monkeypatch):
    template = get_template("srs")
    storage = FakeStorage(fail=True)
    provider = install(monkeypatch, template, storage)
    project, users, _ = await make_project(session, make_user, seed_roles)
    document, _ = await request_document(session, project.id, users["PM"], "srs")
    document_id, project_id, user_id = document.id, project.id, users["PM"].id
    await session.commit()
    await drain(engine)
    session.expire_all()
    document = await session.get(GeneratedDocument, document_id)
    job = await session.get(Job, document.job_id)
    assert document.status == "failed" and document.error_code == "storage_error"
    assert job.status == "failed" and job.error_code == "storage_error"
    assert document.storage_key is None and storage.objects == {}
    calls = len(provider.calls)
    assert calls == len(plan_section_groups(template))

    storage.fail = False
    user = await session.get(User, user_id)
    retried, job = await request_document(session, project_id, user, "srs")
    assert retried.id == document_id and job is not None
    await session.commit()
    await drain(engine)
    session.expire_all()
    document = await session.get(GeneratedDocument, document_id)
    assert document.status == "ready" and document.storage_key in storage.objects
    assert len(provider.calls) == calls

    # Proxy tải: đọc bytes qua service sau khi kiểm tra thành viên.
    await session.refresh(user)
    data, content_type, filename = await document_service.read_document_file(session, document_id, user, storage)
    assert data == storage.objects[document.storage_key][0] and filename == "SRS-v1-vi.docx"


@pytest.mark.asyncio
async def test_missing_items_need_clarification_and_reopen_brief(session, make_user, seed_roles):
    items = json.loads(json.dumps(ITEMS))
    items["key_features"] = {"status": "missing", "value": None, "source": None, "proposed_assumption": None}
    project, users, brief = await make_project(session, make_user, seed_roles, items=items)
    document, job = await request_document(session, project.id, users["BA"], "srs")
    # constraints/out_of_scope skipped không chặn; chỉ mục missing chặn.
    assert job is None and document.status == "needs_clarification"
    assert document.missing_items == ["key_features"]
    assert brief.status == "clarifying" and brief.project_id == project.id
    assert await session.scalar(select(func.count()).select_from(Job)) == 0
    # Yêu cầu lại khi vẫn thiếu: dùng lại bản ghi, không tạo bản ghi mới.
    again, _ = await request_document(session, project.id, users["BA"], "srs")
    assert again.id == document.id


@pytest.mark.asyncio
async def test_no_brief_role_pending_and_quota_rules(session, make_user, seed_roles):
    project, users, _ = await make_project(session, make_user, seed_roles, with_brief=False)
    document, job = await request_document(session, project.id, users["PM"], "brd")
    assert job is None and document.missing_items == ["no_brief"]

    with pytest.raises(HTTPException) as exc:
        await request_document(session, project.id, users["Member"], "srs")
    assert exc.value.status_code == 403
    with pytest.raises(HTTPException) as exc:
        await request_document(session, project.id, users["PM"], "xyz")
    assert (exc.value.status_code, exc.value.detail) == (400, "invalid_doc_type")
    with pytest.raises(HTTPException) as exc:
        await request_document(session, project.id, users["PM"], "srs", language="fr")
    assert (exc.value.status_code, exc.value.detail) == (400, "invalid_language")

    other, other_users, _ = await make_project(session, make_user, seed_roles)
    first, job = await request_document(session, other.id, other_users["BA"], "srs")
    assert job is not None
    with pytest.raises(HTTPException) as exc:
        await request_document(session, other.id, other_users["PM"], "srs")
    assert (exc.value.status_code, exc.value.detail) == (409, "document_generation_pending")
    exhaust_quota(session)
    await session.flush()
    with pytest.raises(HTTPException) as exc:
        await request_document(session, other.id, other_users["PM"], "srs", language="en")
    assert (exc.value.status_code, exc.value.detail) == (429, "ai_quota_exceeded")
