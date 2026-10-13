"""Integration test X2: job `report.xlsx` với storage giả, đối chiếu số liệu với API CPM/dashboard."""
import io
from datetime import date, timedelta

import pytest
import pytest_asyncio
from fastapi import FastAPI
from httpx import ASGITransport, AsyncClient
from openpyxl import load_workbook
from sqlalchemy import insert
from sqlalchemy.ext.asyncio import async_sessionmaker

import app.db.base  # noqa: F401 - đăng ký các quan hệ SQLAlchemy
from app.api.v1.endpoints import reports as reports_endpoint
from app.core.dependencies import get_current_user
from app.core.exceptions import ForbiddenException
from app.core.job_runner import JobRunner
from app.db.session import get_db
from app.models.approval import Approval, ApprovalStatus
from app.models.associations import project_members
from app.models.change_request import ChangeRequest, CRStatus
from app.models.dependency import Dependency
from app.models.job import Job
from app.models.project import Project
from app.models.task import Task, TaskStatus
from app.services.dashboard_service import DashboardService
from app.services.durable_handlers import HANDLERS
from app.services.reports import report_service
from app.services.reports.report_service import compute_eva, request_xlsx_report
from app.services.reports.xlsx_renderer import HEADER_ROW
from app.services.scheduling_service import SchedulingService
from app.services.storage_service import StorageError


class FakeStorage:
    def __init__(self, fail=False):
        self.objects = {}
        self.fail = fail

    async def put_bytes(self, key, data, content_type):
        if self.fail:
            raise StorageError("Supabase storage connection failed")
        self.objects[key] = (data, content_type)

    async def get_bytes(self, key):
        return self.objects[key]


@pytest.fixture
def storage(monkeypatch):
    fake = FakeStorage()
    monkeypatch.setattr(report_service, "get_storage_service", lambda: fake)
    return fake


async def make_project(session, make_user, seed_roles):
    users = {name: await make_user() for name in ("PM", "Member", "Outsider")}
    project = Project(name="Báo cáo XLSX", pm_id=users["PM"].id, start_date=date.today() - timedelta(days=10))
    session.add(project)
    await session.flush()
    for name in ("PM", "Member"):
        await session.execute(insert(project_members).values(
            project_id=project.id, user_id=users[name].id, role_id=seed_roles[name].id,
        ))
    today = date.today()
    a = Task(project_id=project.id, name="A", start_date=today - timedelta(days=10), due_date=today - timedelta(days=6),
             estimated_hours=40, actual_hours=30, progress=100, status=TaskStatus.DONE, actual_end=today - timedelta(days=3))
    b = Task(project_id=project.id, name="B", start_date=today - timedelta(days=5), due_date=today + timedelta(days=4),
             estimated_hours=20, actual_hours=5, progress=50, assignee_id=users["Member"].id)
    c = Task(project_id=project.id, name="=C công thức", start_date=today - timedelta(days=5), due_date=today - timedelta(days=4),
             estimated_hours=10, progress=0)
    session.add_all([a, b, c])
    await session.flush()
    session.add_all([
        Dependency(predecessor_id=a.id, successor_id=b.id),
        Dependency(predecessor_id=a.id, successor_id=c.id),
    ])
    cr = ChangeRequest(project_id=project.id, requested_by_id=users["Member"].id, title="CR một", description="d", status=CRStatus.APPROVED)
    session.add(cr)
    await session.flush()
    session.add(Approval(change_request_id=cr.id, step_order=1, required_role="BA", status=ApprovalStatus.APPROVED,
                         approver_id=users["PM"].id, comment="ok"))
    await session.commit()
    return project, users, [a, b, c]


async def drain(engine, limit=5):
    factory = async_sessionmaker(engine, expire_on_commit=False)
    for _ in range(limit):
        runner = JobRunner(factory, HANDLERS)
        claimed = await runner.poll_once()
        await runner.stop()
        if claimed == 0:
            return


@pytest_asyncio.fixture
async def http_as(session):
    test_app = FastAPI()
    test_app.include_router(reports_endpoint.router, prefix="/api/v1")

    async def override_db():
        yield session
        await session.commit()

    test_app.dependency_overrides[get_db] = override_db

    def factory(user):
        test_app.dependency_overrides[get_current_user] = lambda: user
        return AsyncClient(transport=ASGITransport(app=test_app), base_url="http://test")

    yield factory
    test_app.dependency_overrides.clear()


@pytest.mark.asyncio
async def test_xlsx_matches_cpm_and_dashboard_api(session, engine, make_user, seed_roles, storage, http_as):
    project, users, tasks = await make_project(session, make_user, seed_roles)
    async with http_as(users["Member"]) as client:
        response = await client.post(f"/api/v1/projects/{project.id}/reports/xlsx", json={"language": "vi"})
        assert response.status_code == 202
        job_id = response.json()["job_id"]
        # Job chưa chạy: tải trả 409.
        pending = await client.get(f"/api/v1/projects/{project.id}/reports/xlsx/{job_id}/download")
        assert pending.status_code == 409
        # Chống trùng: yêu cầu lần hai cùng ngôn ngữ trả cùng job.
        again = await client.post(f"/api/v1/projects/{project.id}/reports/xlsx", json={})
        assert again.json()["job_id"] == job_id

    await drain(engine)
    job = await session.get(Job, job_id, populate_existing=True)
    assert job.status == "succeeded", job.error_code
    key = job.result["storage_key"]
    assert key.startswith(f"reports/{project.id}/") and key.endswith(".xlsx")
    assert job.result["file_size"] == len(storage.objects[key][0])
    assert job.result["filename"] == f"project-{project.id}-report-vi.xlsx"

    async with http_as(users["PM"]) as client:
        download = await client.get(f"/api/v1/projects/{project.id}/reports/xlsx/{job_id}/download")
    assert download.status_code == 200
    assert download.headers["cache-control"] == "private, no-store"
    assert download.headers["content-disposition"] == f'attachment; filename="project-{project.id}-report-vi.xlsx"'

    # Số liệu tham chiếu từ chính API CPM và dashboard.
    cpm = await SchedulingService(session).critical_path(project.id, users["PM"])
    stats = await DashboardService(session).get_project_stats(project.id, users["PM"])
    wb = load_workbook(io.BytesIO(download.content))
    task_ws, cp_ws, cr_ws, eva_ws = wb.worksheets

    task_rows = [row for row in task_ws.iter_rows(min_row=HEADER_ROW + 1, values_only=True) if row[0] is not None]
    assert len(task_rows) == stats.total_tasks == len(cpm.tasks) == 3
    # Tên bắt đầu bằng "=" vẫn là chuỗi, không thành công thức.
    assert "=C công thức" in [row[1] for row in task_rows]

    cp_ids = [row[1] for row in cp_ws.iter_rows(min_row=HEADER_ROW + 1, max_row=HEADER_ROW + len(cpm.critical_path), values_only=True)]
    assert cp_ids == cpm.critical_path and cpm.critical_path
    total_row = HEADER_ROW + len(cpm.critical_path) + 2
    assert cp_ws.cell(total_row, 2).value == pytest.approx(cpm.project_duration_days)

    cr_titles = [row[1] for row in cr_ws.iter_rows(min_row=HEADER_ROW + 1, values_only=True) if row[0]]
    assert cr_titles == ["CR một"]

    # Burndown nằm dưới bảng EVA (9 chỉ số) + 1 dòng trống + tiêu đề + hàng tiêu đề.
    bd_header = HEADER_ROW + 9 + 3
    points = list(eva_ws.iter_rows(min_row=bd_header + 1, max_row=bd_header + len(stats.burndown), values_only=True))
    assert [(p[1], p[2]) for p in points] == [(pt.ideal, pt.remaining) for pt in stats.burndown]
    assert eva_ws.cell(HEADER_ROW + 8, 2).value == pytest.approx(70.0)  # BAC


@pytest.mark.asyncio
async def test_outsider_is_rejected(session, make_user, seed_roles, storage, http_as, engine):
    project, users, _ = await make_project(session, make_user, seed_roles)
    async with http_as(users["Outsider"]) as client:
        response = await client.post(f"/api/v1/projects/{project.id}/reports/xlsx", json={})
        assert response.status_code == 403
    job = await request_xlsx_report(session, project.id, users["Member"], "en")
    await session.commit()
    await drain(engine)
    async with http_as(users["Outsider"]) as client:
        download = await client.get(f"/api/v1/projects/{project.id}/reports/xlsx/{job.id}/download")
        assert download.status_code == 403
    # Job thuộc dự án khác: 404.
    other = Project(name="Khác", pm_id=users["Outsider"].id)
    session.add(other)
    await session.commit()
    async with http_as(users["Outsider"]) as client:
        wrong = await client.get(f"/api/v1/projects/{other.id}/reports/xlsx/{job.id}/download")
        assert wrong.status_code == 404
    with pytest.raises(ForbiddenException):
        await request_xlsx_report(session, project.id, users["Outsider"])


@pytest.mark.asyncio
async def test_invalid_language_rejected(session, make_user, seed_roles, storage, http_as):
    project, users, _ = await make_project(session, make_user, seed_roles)
    async with http_as(users["PM"]) as client:
        response = await client.post(f"/api/v1/projects/{project.id}/reports/xlsx", json={"language": "fr"})
    assert response.status_code == 422


@pytest.mark.asyncio
async def test_storage_error_fails_job_and_retry_works(session, engine, make_user, seed_roles, storage, monkeypatch):
    project, users, _ = await make_project(session, make_user, seed_roles)
    storage.fail = True
    job = await request_xlsx_report(session, project.id, users["PM"])
    await session.commit()
    await drain(engine)
    failed = await session.get(Job, job.id, populate_existing=True)
    assert failed.status == "failed" and failed.error_code == "storage_error"
    with pytest.raises(Exception) as exc:
        await report_service.read_report_file(session, project.id, job.id, users["PM"])
    assert exc.value.status_code == 409

    storage.fail = False
    retry = await request_xlsx_report(session, project.id, users["PM"])
    assert retry.id != job.id
    await session.commit()
    await drain(engine)
    assert (await session.get(Job, retry.id, populate_existing=True)).status == "succeeded"


@pytest.mark.asyncio
async def test_member_removed_before_run_fails_job(session, engine, make_user, seed_roles, storage):
    project, users, _ = await make_project(session, make_user, seed_roles)
    job = await request_xlsx_report(session, project.id, users["Member"])
    await session.execute(project_members.delete().where(project_members.c.user_id == users["Member"].id))
    await session.commit()
    await drain(engine)
    failed = await session.get(Job, job.id, populate_existing=True)
    assert failed.status == "failed" and failed.error_code == "project_unavailable"
    assert storage.objects == {}


def test_compute_eva_formula():
    today = date(2026, 1, 11)
    tasks = [
        Task(name="x", estimated_hours=10, actual_hours=4, progress=50, start_date=date(2026, 1, 1), due_date=date(2026, 1, 20)),
        Task(name="y", estimated_hours=None, actual_hours=2, progress=0),
    ]
    eva = compute_eva(tasks, today)
    assert eva.bac == 10 and eva.ev == 5 and eva.ac == 6
    assert eva.pv == pytest.approx(5.5)  # 11/20 thời gian đã trôi qua
    assert eva.cpi == pytest.approx(round(5 / 6, 3))
    assert compute_eva([Task(name="z", estimated_hours=None, actual_hours=0, progress=0)], today) is None


def test_report_tasks_celery_module_removed():
    import importlib.util

    assert importlib.util.find_spec("app.workers.report_tasks") is None
