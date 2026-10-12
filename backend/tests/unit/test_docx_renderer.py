"""D1: mẫu SRS/BRD và bộ dựng DOCX thuần.

Đọc lại DOCX bằng python-docx để kiểm tra: tiêu đề đúng cấp, mục lục (field TOC),
bảng yêu cầu chức năng, mục giả định sinh từ hồ sơ, loại key lạ, ngôn ngữ vi/en,
chèn ảnh PNG và bảng bước dự phòng.
"""
import io
import struct
import zlib
from datetime import datetime

import pytest
from docx import Document

from app.services.documents import (
    BRD_TEMPLATE,
    SRS_TEMPLATE,
    DocumentRenderError,
    get_template,
    missing_brief_items,
    missing_required_sections,
    render_docx,
    render_docx_report,
    validate_content,
)

GENERATED_AT = datetime(2026, 10, 12, 9, 30)

FUNCTIONAL_ROWS = [
    {
        "id": "FR-01",
        "name": "Đăng nhập",
        "description": "Người dùng đăng nhập bằng email",
        "priority": "Cao",
        "acceptance_criteria": ["Sai mật khẩu báo lỗi", "Khóa sau 5 lần"],
    },
    {
        "id": "FR-02",
        "name": "Tạo dự án",
        "description": "PM tạo dự án mới",
        "priority": "Trung bình",
        "acceptance_criteria": "Dự án xuất hiện trong danh sách",
    },
]

ASSUMPTIONS = [
    {"key": "constraints", "label": "Ràng buộc", "status": "assumed", "value": "Ngân sách 500 triệu"},
    {"key": "success_criteria", "label": "Tiêu chí thành công", "status": "skipped", "value": None},
    {"key": "goals", "label": "Mục tiêu", "status": "clear", "value": "Không được hiện"},
]


def _png_bytes(width: int = 2, height: int = 2) -> bytes:
    """Tạo ảnh PNG RGB nhỏ hợp lệ tại chỗ (không cần thư viện ảnh)."""

    def chunk(kind: bytes, data: bytes) -> bytes:
        body = kind + data
        return struct.pack(">I", len(data)) + body + struct.pack(">I", zlib.crc32(body) & 0xFFFFFFFF)

    raw = b"".join(b"\x00" + b"\xff\x00\x00" * width for _ in range(height))
    header = struct.pack(">IIBBBBB", width, height, 8, 2, 0, 0, 0)
    return (
        b"\x89PNG\r\n\x1a\n"
        + chunk(b"IHDR", header)
        + chunk(b"IDAT", zlib.compress(raw))
        + chunk(b"IEND", b"")
    )


def _render(template=SRS_TEMPLATE, content=None, assumptions=None, **kwargs):
    kwargs.setdefault("title", "Cổng khách hàng")
    kwargs.setdefault("project_name", "Portal Revamp")
    kwargs.setdefault("generated_at", GENERATED_AT)
    return render_docx_report(
        template,
        content if content is not None else {"functional_requirements": FUNCTIONAL_ROWS},
        assumptions if assumptions is not None else ASSUMPTIONS,
        **kwargs,
    )


def _open(data: bytes):
    return Document(io.BytesIO(data))


def _headings(document) -> list[tuple[int, str]]:
    result = []
    for paragraph in document.paragraphs:
        name = paragraph.style.name
        if name.startswith("Heading "):
            result.append((int(name.split()[1]), paragraph.text))
    return result


def _table_after_heading(document, heading_text: str):
    """Tìm bảng đứng ngay sau tiêu đề có chứa ``heading_text`` trong thân tài liệu."""
    body = document.element.body
    found = False
    for element in body.iterchildren():
        tag = element.tag.rsplit("}", 1)[-1]
        if tag == "p":
            text = "".join(node.text or "" for node in element.iter() if node.tag.endswith("}t"))
            style = element.find(".//{*}pStyle")
            is_heading = style is not None and "Heading" in (style.get(
                "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val") or "")
            if is_heading:
                if found:
                    return None
                found = heading_text in text
        elif tag == "tbl" and found:
            for table in document.tables:
                if table._tbl is element:
                    return table
    return None


def _rows(table) -> list[list[str]]:
    return [[cell.text for cell in row.cells] for row in table.rows]


# ---------------------------------------------------------------------------
# Mẫu và ánh xạ hồ sơ
# ---------------------------------------------------------------------------


def test_templates_have_unique_keys_and_valid_structure():
    for template in (SRS_TEMPLATE, BRD_TEMPLATE):
        keys = template.section_keys
        assert len(keys) == len(set(keys))
        assert keys[-1] == "assumptions_open_issues"
        assert "business_process" in keys
        for section in template.sections:
            assert section.level in (1, 2)
            assert section.title_vi and section.title_en
            if section.kind == "table":
                assert section.columns
            else:
                assert not section.columns
    fr = SRS_TEMPLATE.section("functional_requirements")
    assert fr.kind == "table" and fr.required
    assert [c.key for c in fr.columns] == ["id", "name", "description", "priority", "acceptance_criteria"]
    assert get_template("BRD") is BRD_TEMPLATE
    with pytest.raises(ValueError):
        get_template("pdf")


def test_missing_brief_items_counts_only_clear_and_assumed():
    brief = {
        "goals": {"status": "clear", "value": "Giảm thời gian xử lý"},
        "stakeholders": {"status": "assumed", "value": "PM, BA"},
        "in_scope": {"status": "skipped", "value": None},
        "key_features": {"status": "missing", "value": None},
        "constraints": {"status": "clear", "value": "Ngân sách"},
    }
    # SRS cần goals, in_scope, key_features, stakeholders, constraints.
    assert missing_brief_items(SRS_TEMPLATE, brief) == ["in_scope", "key_features"]
    # BRD cần thêm success_criteria; out_of_scope không bắt buộc.
    assert missing_brief_items(BRD_TEMPLATE, brief) == ["in_scope", "key_features", "success_criteria"]
    full = {key: {"status": "clear", "value": "x"} for key in
            ("goals", "stakeholders", "in_scope", "key_features", "constraints", "success_criteria")}
    assert missing_brief_items(BRD_TEMPLATE, full) == []
    assert "out_of_scope" not in missing_brief_items(BRD_TEMPLATE, {})


# ---------------------------------------------------------------------------
# Dựng DOCX
# ---------------------------------------------------------------------------


def test_docx_opens_with_title_toc_and_heading_levels():
    result = _render()
    document = _open(result.data)

    assert document.paragraphs[0].style.name == "Title"
    assert document.paragraphs[0].text == "Cổng khách hàng"
    assert document.core_properties.title == "Cổng khách hàng"
    full_text = "\n".join(p.text for p in document.paragraphs)
    assert "Portal Revamp" in full_text and "12/10/2026 09:30" in full_text

    xml = document.element.xml
    assert 'TOC \\o "1-2"' in xml
    assert "Mục lục" in full_text
    assert document.settings.element.find(".//{*}updateFields") is not None

    headings = _headings(document)
    assert headings[0] == (1, "1. Giới thiệu")
    assert (2, "1.1. Mục đích") in headings
    assert (2, "1.3. Định nghĩa thuật ngữ") in headings
    assert (1, "3. Yêu cầu chức năng") in headings
    assert headings[-1] == (1, "7. Giả định và vấn đề mở")
    assert len(headings) == len(SRS_TEMPLATE.sections)


def test_functional_requirements_table_has_header_and_rows():
    document = _open(_render().data)
    table = _table_after_heading(document, "Yêu cầu chức năng")
    assert table is not None
    rows = _rows(table)
    assert rows[0] == ["Mã", "Tên", "Mô tả", "Ưu tiên", "Tiêu chí chấp nhận"]
    assert rows[1][:4] == ["FR-01", "Đăng nhập", "Người dùng đăng nhập bằng email", "Cao"]
    assert rows[1][4] == "Sai mật khẩu báo lỗi\nKhóa sau 5 lần"
    assert rows[2][0] == "FR-02"
    # Hàng tiêu đề được đánh dấu lặp lại khi sang trang.
    assert table.rows[0]._tr.find(".//{*}tblHeader") is not None


def test_assumptions_section_lists_assumed_and_skipped_from_brief_only():
    content = {
        "functional_requirements": FUNCTIONAL_ROWS,
        # AI cố tự viết mục giả định: phải bị loại.
        "assumptions_open_issues": "AI tự bịa giả định",
    }
    result = _render(content=content)
    document = _open(result.data)
    table = _table_after_heading(document, "Giả định và vấn đề mở")
    rows = _rows(table)
    assert rows[0] == ["Mục", "Trạng thái", "Nội dung"]
    assert rows[1] == ["Ràng buộc", "Giả định", "Ngân sách 500 triệu"]
    assert rows[2] == ["Tiêu chí thành công", "Vấn đề mở", "Chưa có thông tin, cần làm rõ."]
    assert len(rows) == 3  # mục "clear" không vào bảng giả định
    full = "\n".join(p.text for p in document.paragraphs)
    assert "AI tự bịa giả định" not in full
    assert any("assumptions_open_issues" in w for w in result.warnings)
    assert any("'clear'" in w for w in result.warnings)


def test_no_assumptions_prints_note():
    document = _open(_render(assumptions=[]).data)
    full = "\n".join(p.text for p in document.paragraphs)
    assert "Không có giả định hay vấn đề mở." in full
    assert _table_after_heading(document, "Giả định và vấn đề mở") is None


def test_unknown_keys_and_wrong_types_are_dropped_with_warnings():
    content = {
        "functional_requirements": FUNCTIONAL_ROWS,
        "budget_secret": "Không thuộc mẫu",
        "purpose": ["Đoạn một", "Đoạn hai"],
        "constraints": "phải là danh sách",  # sai kiểu: bullets cần list[str]
        "introduction": "mục nhóm không nhận nội dung",
        "stakeholders": [{"role": "PM", "description": "Quản lý", "salary": 1000}],
    }
    result = _render(content=content)
    document = _open(result.data)
    full = "\n".join(p.text for p in document.paragraphs)
    assert "Không thuộc mẫu" not in full
    assert "phải là danh sách" not in full
    assert "Đoạn một" in full and "Đoạn hai" in full
    joined = " ".join(result.warnings)
    assert "budget_secret" in joined
    assert "'constraints'" in joined
    assert "'introduction'" in joined
    assert "salary" in joined
    stakeholders = _rows(_table_after_heading(document, "Bên liên quan"))
    assert stakeholders == [["Vai trò", "Mô tả"], ["PM", "Quản lý"]]

    cleaned = validate_content(SRS_TEMPLATE, content)
    assert set(cleaned.content) == {"functional_requirements", "purpose", "stakeholders"}


def test_table_with_nested_object_cell_is_rejected():
    content = {"functional_requirements": [{"id": "FR-1", "name": {"x": 1}}]}
    cleaned = validate_content(SRS_TEMPLATE, content)
    assert "functional_requirements" not in cleaned.content
    assert missing_required_sections(SRS_TEMPLATE, content)[0] == "purpose"
    assert "functional_requirements" in missing_required_sections(SRS_TEMPLATE, content)


def test_missing_required_sections_show_no_information_note():
    document = _open(_render(content={}).data)
    paragraphs = document.paragraphs
    for index, paragraph in enumerate(paragraphs):
        if paragraph.text == "1.1. Mục đích":
            assert paragraphs[index + 1].text == "Chưa có thông tin."
            assert paragraphs[index + 1].runs[0].italic
            break
    else:
        pytest.fail("Không thấy tiêu đề Mục đích")
    assert _table_after_heading(document, "Yêu cầu chức năng") is None


def test_english_language_uses_english_titles():
    result = _render(template=BRD_TEMPLATE, content={"in_scope": ["Portal"]}, language="en")
    document = _open(result.data)
    headings = _headings(document)
    assert headings[0] == (1, "1. Executive Summary")
    assert (1, "5. Scope") in headings
    assert (2, "5.1. In Scope") in headings
    assert headings[-1][1].endswith("Assumptions and Open Issues")
    full = "\n".join(p.text for p in document.paragraphs)
    assert "Table of Contents" in full
    assert "No information available yet." in full
    assert "2026-10-12 09:30" in full
    rows = _rows(_table_after_heading(document, "Assumptions and Open Issues"))
    assert rows[0] == ["Item", "Status", "Details"]
    assert rows[1][1] == "Assumption"
    assert rows[2][1] == "Open issue"
    bullets = [p.text for p in document.paragraphs if p.style.name == "List Bullet"]
    assert bullets == ["Portal"]


def test_unsupported_language_raises():
    with pytest.raises(DocumentRenderError):
        _render(language="fr")


def test_process_diagram_png_is_embedded():
    png = _png_bytes()
    steps = [{"order": 1, "lane": "BA", "name": "Thu thập yêu cầu"}]
    result = _render(process_diagram_png=png, process_steps=steps)
    document = _open(result.data)
    assert len(document.inline_shapes) == 1
    images = [rel for rel in document.part.rels.values() if "image" in rel.reltype]
    assert len(images) == 1
    assert images[0].target_part.blob == png
    # Có ảnh thì không in bảng bước dự phòng.
    assert _table_after_heading(document, "Quy trình nghiệp vụ") is None
    assert not any("sơ đồ" in w for w in result.warnings)


def test_process_steps_table_fallback_without_diagram():
    steps = [
        {"order": 1, "lane": "BA", "name": "Thu thập yêu cầu", "description": "Phỏng vấn"},
        {"lane": "PO", "name": "Duyệt phạm vi"},
        {"lane": "PM"},  # không có tên/mô tả: bỏ
    ]
    document = _open(_render(process_steps=steps).data)
    assert len(document.inline_shapes) == 0
    rows = _rows(_table_after_heading(document, "Quy trình nghiệp vụ"))
    assert rows[0] == ["Bước", "Vai trò", "Công việc", "Mô tả"]
    assert rows[1] == ["1", "BA", "Thu thập yêu cầu", "Phỏng vấn"]
    assert rows[2] == ["2", "PO", "Duyệt phạm vi", ""]
    assert len(rows) == 3


def test_invalid_png_falls_back_to_steps_table_with_warning():
    steps = [{"order": 1, "lane": "BA", "name": "Thu thập yêu cầu"}]
    result = _render(process_diagram_png=b"not a png", process_steps=steps)
    document = _open(result.data)
    assert len(document.inline_shapes) == 0
    assert _table_after_heading(document, "Quy trình nghiệp vụ") is not None
    assert any("sơ đồ" in w for w in result.warnings)


def test_render_docx_returns_bytes_of_valid_docx():
    data = render_docx(
        BRD_TEMPLATE,
        {"business_requirements": [{"id": "BR-1", "name": "Báo cáo", "description": "Xuất XLSX",
                                    "priority": "Cao"}]},
        ASSUMPTIONS,
        title="BRD",
        project_name="Demo",
        generated_at=GENERATED_AT,
    )
    assert data[:2] == b"PK"
    document = _open(data)
    rows = _rows(_table_after_heading(document, "Yêu cầu kinh doanh"))
    assert rows[0] == ["Mã", "Tên", "Mô tả", "Ưu tiên"]
    assert rows[1] == ["BR-1", "Báo cáo", "Xuất XLSX", "Cao"]
