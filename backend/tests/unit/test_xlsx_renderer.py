"""Test renderer XLSX thuần: sinh tệp rồi đọc lại bằng openpyxl để kiểm tra."""
from datetime import UTC, date, datetime
from io import BytesIO
from typing import Any

import pytest
from openpyxl import load_workbook

from app.schemas.dashboard import BurndownPoint
from app.services.reports import (
    BurndownRow,
    ChangeRequestRow,
    CriticalPathData,
    EVAData,
    ProjectInfo,
    ProjectReportData,
    ReportApprovalStep,
    TaskRow,
    render_project_xlsx,
    sheet_names,
)
from app.services.reports.xlsx_renderer import HEADER_ROW, burndown_header_row


def _full_data() -> ProjectReportData:
    tasks = [
        TaskRow(
            id=1, name="Phân tích yêu cầu", phase="Khởi tạo", status="DONE",
            priority="HIGH", assignee="An", start_date=date(2026, 1, 5),
            due_date=date(2026, 1, 9), progress=100.0, estimated_hours=16.0,
            actual_hours=18.5, duration_days=5.0, early_start=date(2026, 1, 5),
            early_finish=date(2026, 1, 9), late_start=date(2026, 1, 5),
            late_finish=date(2026, 1, 9), float_days=0.0, is_critical=True,
        ),
        TaskRow(
            id=2, name="Viết tài liệu", sprint="Sprint 1", epic="Docs",
            status="IN_PROGRESS", priority="LOW", start_date=date(2026, 1, 12),
            due_date=date(2026, 1, 14), progress=40.0, estimated_hours=8.0,
            actual_hours=3.0, duration_days=3.0, float_days=4.0, is_critical=False,
        ),
        TaskRow(
            id=3, name="Triển khai", status="TODO", priority="MEDIUM",
            duration_days=2.0, early_start=date(2026, 1, 12),
            early_finish=date(2026, 1, 13), float_days=0.0, is_critical=True,
        ),
    ]
    crs = [
        ChangeRequestRow(
            id=7, title="Thêm module báo cáo", status="APPROVED",
            requested_by="Bình", created_at=datetime(2026, 2, 1, 9, 30, tzinfo=UTC),
            approvals=[
                ReportApprovalStep(step_order=2, decision="APPROVED", decided_by="PO X",
                                   decided_at=datetime(2026, 2, 3, 10, 0)),
                ReportApprovalStep(step_order=1, decision="APPROVED", decided_by="BA Y",
                                   decided_at=datetime(2026, 2, 2, 8, 0), reason="Hợp lý"),
                ReportApprovalStep(step_order=3, decision="APPROVED", decided_by="PM Z",
                                   decided_at=datetime(2026, 2, 4, 15, 0)),
            ],
        ),
        ChangeRequestRow(id=8, title="Đổi màu", status="DRAFT", requested_by="Cường",
                         created_at=date(2026, 2, 5)),
    ]
    return ProjectReportData(
        project=ProjectInfo(name="Dự án Demo", code="DEMO", exported_at=date(2026, 3, 1)),
        tasks=tasks,
        critical_path=CriticalPathData(task_ids=[1, 3], project_duration_days=9.0),
        change_requests=crs,
        eva=EVAData(pv=1000.0, ev=800.0, ac=900.0, spi=0.8, cpi=0.888, sv=-200.0,
                    cv=-100.0, bac=5000.0, eac=None),
        burndown=[
            BurndownRow.from_dashboard_point(
                BurndownPoint(date="2026-02-01", remaining=3.0, ideal=3.0)),
            BurndownRow(date=date(2026, 2, 2), ideal=1.5, actual=2.0),
        ],
    )


def _load(data: ProjectReportData, language: str = "vi"):
    return load_workbook(BytesIO(render_project_xlsx(data, language=language)))


def _row(ws, row: int) -> list:
    return [c.value for c in ws[row]]


def test_ten_sheet_theo_ngon_ngu():
    assert _load(_full_data(), "vi").sheetnames == [
        "Task", "Đường găng", "CR và duyệt", "EVA và Burndown"]
    assert _load(_full_data(), "en").sheetnames == [
        "Task", "Critical path", "Change requests", "EVA & Burndown"]
    assert list(sheet_names("en")) == _load(_full_data(), "en").sheetnames


def test_ngon_ngu_khong_ho_tro_bao_loi():
    with pytest.raises(ValueError):
        unsupported: Any = "fr"
        render_project_xlsx(_full_data(), language=unsupported)


def test_sheet_task_tieu_de_kieu_du_lieu_va_cot_gang():
    wb = _load(_full_data())
    ws = wb["Task"]
    headers = _row(ws, HEADER_ROW)
    assert headers[:3] == ["ID", "Tên task", "Phase"]
    assert headers[-1] == "Găng"
    assert all(c.font.bold for c in ws[HEADER_ROW])
    assert ws.freeze_panes == f"A{HEADER_ROW + 1}"
    # Thông tin dự án ở đầu sheet
    assert ws["B1"].value == "Dự án Demo (DEMO)"
    assert ws["B2"].is_date
    # 3 task => 3 hàng dữ liệu
    assert ws.max_row == HEADER_ROW + 3
    first = ws[HEADER_ROW + 1]
    col = {h: i for i, h in enumerate(headers)}
    assert first[col["Bắt đầu"]].is_date
    assert first[col["Bắt đầu"]].value.date() == date(2026, 1, 5)
    assert first[col["Tiến độ (%)"]].data_type == "n"
    assert first[col["Tiến độ (%)"]].value == 100.0
    assert first[col["Giờ thực tế"]].value == 18.5
    assert isinstance(first[col["ID"]].value, int)
    gang = [ws.cell(row=r, column=col["Găng"] + 1).value
            for r in range(HEADER_ROW + 1, HEADER_ROW + 4)]
    assert gang == ["Có", "Không", "Có"]
    # Ô trống khi không có dữ liệu (không ghi chuỗi "None")
    assert ws.cell(row=HEADER_ROW + 3, column=col["Bắt đầu"] + 1).value is None
    # Độ rộng cột hợp lý
    width = ws.column_dimensions["B"].width
    assert 8 <= width <= 50


def test_sheet_task_tieng_anh_dung_yes_no():
    ws = _load(_full_data(), "en")["Task"]
    headers = _row(ws, HEADER_ROW)
    assert headers[-1] == "Critical"
    assert "Task name" in headers
    idx = headers.index("Critical") + 1
    assert [ws.cell(row=r, column=idx).value for r in range(HEADER_ROW + 1, HEADER_ROW + 4)] == [
        "Yes", "No", "Yes"]


def test_sheet_duong_gang_theo_thu_tu_va_tong_thoi_luong():
    ws = _load(_full_data())["Đường găng"]
    assert _row(ws, HEADER_ROW)[:4] == ["Thứ tự", "ID", "Tên task", "Thời lượng (ngày)"]
    assert _row(ws, HEADER_ROW + 1)[:4] == [1, 1, "Phân tích yêu cầu", 5.0]
    assert _row(ws, HEADER_ROW + 2)[:4] == [2, 3, "Triển khai", 2.0]
    assert ws.cell(row=HEADER_ROW + 1, column=5).is_date
    assert ws.cell(row=HEADER_ROW + 1, column=10).value == "Có"
    total_row = HEADER_ROW + 2 + 2
    assert ws.cell(row=total_row, column=1).value == "Tổng thời lượng dự án (ngày)"
    assert ws.cell(row=total_row, column=2).value == 9.0
    assert ws.freeze_panes == f"A{HEADER_ROW + 1}"


def test_sheet_cr_moi_buoc_duyet_mot_hang_va_vai_tro_mac_dinh():
    ws = _load(_full_data())["CR và duyệt"]
    headers = _row(ws, HEADER_ROW)
    assert headers == ["Mã CR", "Tiêu đề", "Trạng thái CR", "Người tạo", "Ngày tạo",
                       "Bước", "Vai trò", "Quyết định", "Người quyết định", "Thời điểm",
                       "Lý do"]
    # CR-7 có 3 bước, CR-8 không có bước nào => 4 hàng
    assert ws.max_row == HEADER_ROW + 4
    rows = [_row(ws, r) for r in range(HEADER_ROW + 1, HEADER_ROW + 5)]
    assert [r[0] for r in rows] == ["CR-7", "CR-7", "CR-7", "CR-8"]
    assert [r[5] for r in rows[:3]] == [1, 2, 3]  # đã sắp theo thứ tự bước
    assert [r[6] for r in rows[:3]] == ["BA", "PO", "PM"]
    assert rows[0][10] == "Hợp lý"
    # datetime có tz được quy về UTC naive và vẫn là ngày thật
    created = ws.cell(row=HEADER_ROW + 1, column=5)
    assert created.is_date and created.value == datetime(2026, 2, 1, 9, 30)
    assert ws.cell(row=HEADER_ROW + 1, column=10).is_date
    assert rows[3][5] is None and rows[3][7] is None


def test_sheet_eva_va_burndown():
    data = _full_data()
    ws = _load(data, "en")["EVA & Burndown"]
    assert _row(ws, HEADER_ROW)[:2] == ["Metric", "Value"]
    metrics = {ws.cell(row=r, column=1).value: ws.cell(row=r, column=2)
               for r in range(HEADER_ROW + 1, HEADER_ROW + 10)}
    assert metrics["PV - Planned value"].value == 1000.0
    assert metrics["PV - Planned value"].data_type == "n"
    assert metrics["SV - Schedule variance"].value == -200.0
    assert metrics["EAC - Estimate at completion"].value is None
    bd = burndown_header_row(data)
    assert _row(ws, bd)[:3] == ["Date", "Ideal", "Actual (remaining)"]
    assert ws.cell(row=bd, column=1).font.bold
    first = ws.cell(row=bd + 1, column=1)
    assert first.is_date and first.value.date() == date(2026, 2, 1)
    assert ws.cell(row=bd + 1, column=3).value == 3.0
    assert ws.cell(row=bd + 2, column=2).value == 1.5
    assert ws.max_row == bd + 2


def test_danh_sach_rong_van_co_tieu_de_va_dong_khong_co_du_lieu():
    data = ProjectReportData(project=ProjectInfo(name="Rỗng"))
    wb = _load(data)
    for name in wb.sheetnames:
        ws = wb[name]
        assert ws.cell(row=HEADER_ROW, column=1).font.bold
        assert ws.cell(row=HEADER_ROW + 1, column=1).value == "Không có dữ liệu"
    ws = wb["EVA và Burndown"]
    bd = burndown_header_row(data)
    assert ws.cell(row=bd, column=1).value == "Ngày"
    assert ws.cell(row=bd + 1, column=1).value == "Không có dữ liệu"
    # Ngày xuất mặc định là hôm nay, dạng ngày thật
    assert wb["Task"]["B2"].is_date

    en = _load(data, "en")
    assert en["Task"].cell(row=HEADER_ROW + 1, column=1).value == "No data"


@pytest.mark.parametrize("payload", ["=1+1", "+cmd|' /C calc'!A0", "-2+3", "@SUM(A1)",
                                     "=HYPERLINK(\"http://x\",\"y\")"])
def test_chong_formula_injection(payload):
    data = ProjectReportData(
        project=ProjectInfo(name=payload, code=None),
        tasks=[TaskRow(id=1, name=payload, assignee=payload, status="TODO")],
        critical_path=CriticalPathData(task_ids=[1], project_duration_days=1.0),
        change_requests=[ChangeRequestRow(
            id=1, title=payload, requested_by=payload,
            approvals=[ReportApprovalStep(step_order=1, decision="REJECTED", reason=payload)])],
    )
    raw = render_project_xlsx(data)
    wb = load_workbook(BytesIO(raw))
    cells = [
        wb["Task"]["B1"],
        wb["Task"].cell(row=HEADER_ROW + 1, column=2),
        wb["Task"].cell(row=HEADER_ROW + 1, column=8),
        wb["Đường găng"].cell(row=HEADER_ROW + 1, column=3),
        wb["CR và duyệt"].cell(row=HEADER_ROW + 1, column=2),
        wb["CR và duyệt"].cell(row=HEADER_ROW + 1, column=11),
    ]
    for cell in cells:
        assert cell.value == payload
        assert cell.data_type == "s"
    # Không có thẻ công thức <f> nào trong XML của các sheet
    import zipfile
    with zipfile.ZipFile(BytesIO(raw)) as zf:
        for name in zf.namelist():
            if name.startswith("xl/worksheets/sheet"):
                assert b"<f>" not in zf.read(name)


def test_ky_tu_dieu_khien_bi_loai_bo():
    data = ProjectReportData(
        project=ProjectInfo(name="P"),
        tasks=[TaskRow(id=1, name="ab\x00c\x07d")],
    )
    ws = _load(data)["Task"]
    assert ws.cell(row=HEADER_ROW + 1, column=2).value == "abcd"
