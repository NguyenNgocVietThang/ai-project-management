"""Dựng tài liệu DOCX (SRS/BRD) từ mẫu cố định và nội dung JSON có cấu trúc.

Hàm thuần, không gọi AI, không truy cập DB hay storage. Nội dung AI chỉ được điền
vào các mục của mẫu; key lạ hoặc sai kiểu bị loại kèm cảnh báo (không raise), để
một mục hỏng không làm hỏng cả tài liệu. Mục "Giả định và vấn đề mở" luôn sinh từ
tham số ``assumptions`` (dữ liệu hồ sơ), không lấy từ nội dung AI.
"""
from __future__ import annotations

import io
import logging
from collections.abc import Mapping, Sequence
from dataclasses import dataclass, field
from datetime import datetime
from typing import Any

from docx import Document
from docx.document import Document as DocxDocument
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

from app.services.documents.templates import (
    DEFAULT_LANGUAGE,
    SUPPORTED_LANGUAGES,
    DocumentTemplate,
    Section,
)

logger = logging.getLogger(__name__)

# Kiểu giá trị đơn được phép trong ô bảng.
_SCALAR_TYPES = (str, int, float, bool)

# Trạng thái hợp lệ của phần tử ``assumptions``.
ASSUMPTION_STATUSES: frozenset[str] = frozenset({"assumed", "skipped"})

# Nhãn cố định theo ngôn ngữ.
_LABELS: dict[str, dict[str, str]] = {
    "vi": {
        "toc": "Mục lục",
        "toc_placeholder": (
            "Nhấn chuột phải vào đây và chọn \"Cập nhật trường\" (F9) để hiển thị mục lục."
        ),
        "project": "Dự án",
        "generated_at": "Ngày tạo",
        "no_information": "Chưa có thông tin.",
        "no_assumptions": "Không có giả định hay vấn đề mở.",
        "assumption_item": "Mục",
        "assumption_status": "Trạng thái",
        "assumption_value": "Nội dung",
        "status_assumed": "Giả định",
        "status_skipped": "Vấn đề mở",
        "skipped_value": "Chưa có thông tin, cần làm rõ.",
        "diagram_caption": "Sơ đồ quy trình nghiệp vụ",
        "step_order": "Bước",
        "step_lane": "Vai trò",
        "step_name": "Công việc",
        "step_description": "Mô tả",
    },
    "en": {
        "toc": "Table of Contents",
        "toc_placeholder": (
            "Right-click here and choose \"Update Field\" (F9) to display the table of contents."
        ),
        "project": "Project",
        "generated_at": "Generated at",
        "no_information": "No information available yet.",
        "no_assumptions": "No assumptions or open issues.",
        "assumption_item": "Item",
        "assumption_status": "Status",
        "assumption_value": "Details",
        "status_assumed": "Assumption",
        "status_skipped": "Open issue",
        "skipped_value": "No information available yet; needs clarification.",
        "diagram_caption": "Business process diagram",
        "step_order": "Step",
        "step_lane": "Role",
        "step_name": "Activity",
        "step_description": "Description",
    },
}

# Cột của bảng bước quy trình dự phòng: key trong ``process_steps`` -> key nhãn.
PROCESS_STEP_COLUMNS: tuple[tuple[str, str], ...] = (
    ("order", "step_order"),
    ("lane", "step_lane"),
    ("name", "step_name"),
    ("description", "step_description"),
)


class DocumentRenderError(ValueError):
    """Tham số dựng tài liệu không hợp lệ (ngôn ngữ, mẫu...)."""


@dataclass
class ContentValidation:
    """Kết quả kiểm tra nội dung AI: nội dung đã chuẩn hóa và danh sách cảnh báo."""

    content: dict[str, Any] = field(default_factory=dict)
    warnings: list[str] = field(default_factory=list)


@dataclass
class RenderResult:
    """Kết quả dựng DOCX: bytes của tệp và các cảnh báo trong quá trình dựng."""

    data: bytes
    warnings: list[str] = field(default_factory=list)


def labels(language: str) -> dict[str, str]:
    """Nhãn cố định theo ngôn ngữ (ngôn ngữ lạ gây DocumentRenderError)."""
    _ensure_language(language)
    return _LABELS[language]


def _ensure_language(language: str) -> None:
    if language not in SUPPORTED_LANGUAGES:
        raise DocumentRenderError(f"Ngôn ngữ không hỗ trợ: {language!r}")


def _clean_text(value: Any) -> str | None:
    """Chuỗi đã cắt khoảng trắng; chuỗi rỗng coi như không có."""
    if not isinstance(value, str):
        return None
    text = value.strip()
    return text or None


def _cell_text(value: Any) -> str | None:
    """Chuyển giá trị ô bảng thành chuỗi; trả None nếu sai kiểu."""
    if value is None:
        return ""
    if isinstance(value, _SCALAR_TYPES):
        return str(value).strip()
    if isinstance(value, Sequence) and not isinstance(value, str | bytes):
        parts: list[str] = []
        for item in value:
            if not isinstance(item, _SCALAR_TYPES):
                return None
            text = str(item).strip()
            if text:
                parts.append(text)
        return "\n".join(parts)
    return None


def _normalize_paragraph(value: Any) -> list[str] | None:
    if isinstance(value, str):
        text = _clean_text(value)
        return [text] if text else []
    if isinstance(value, Sequence) and not isinstance(value, bytes):
        if not all(isinstance(item, str) for item in value):
            return None
        return [text for text in (_clean_text(item) for item in value) if text]
    return None


def _normalize_bullets(value: Any) -> list[str] | None:
    if isinstance(value, str) or not isinstance(value, Sequence):
        return None
    if not all(isinstance(item, str) for item in value):
        return None
    return [text for text in (_clean_text(item) for item in value) if text]


def _normalize_table(
    section: Section, value: Any, warnings: list[str]
) -> list[dict[str, str]] | None:
    if isinstance(value, str | bytes) or not isinstance(value, Sequence):
        return None
    column_keys = [column.key for column in section.columns]
    rows: list[dict[str, str]] = []
    for index, row in enumerate(value):
        if not isinstance(row, Mapping):
            return None
        extra = sorted(str(key) for key in row if key not in column_keys)
        if extra:
            warnings.append(
                f"Mục '{section.key}' hàng {index + 1}: bỏ cột không thuộc mẫu {extra}."
            )
        normalized: dict[str, str] = {}
        for key in column_keys:
            text = _cell_text(row.get(key))
            if text is None:
                return None
            normalized[key] = text
        # Bỏ hàng rỗng hoàn toàn để không in hàng trắng.
        if any(normalized.values()):
            rows.append(normalized)
    return rows


def validate_content(
    template: DocumentTemplate, content: Mapping[str, Any] | None
) -> ContentValidation:
    """Kiểm tra và chuẩn hóa nội dung AI theo mẫu.

    Quy tắc: key không thuộc mẫu bị loại; mục ``group``/``assumptions`` không nhận
    nội dung từ AI; giá trị sai kiểu bị loại cả mục. Mọi lần loại đều ghi cảnh báo,
    không raise. Mục còn lại được chuẩn hóa:

    - paragraph: list[str] (mỗi phần tử một đoạn văn)
    - bullets: list[str]
    - table: list[dict[key cột, str]]
    - process: list[str] (đoạn mô tả tùy chọn)
    """
    result = ContentValidation()
    if content is None:
        return result
    if not isinstance(content, Mapping):
        result.warnings.append("Nội dung tài liệu không phải đối tượng JSON; bỏ toàn bộ.")
        return result

    for raw_key, value in content.items():
        key = str(raw_key)
        section = template.section(key)
        if section is None:
            result.warnings.append(f"Bỏ mục không thuộc mẫu: '{key}'.")
            continue
        if not section.accepts_ai_content:
            result.warnings.append(
                f"Bỏ nội dung của mục '{key}': mục này không nhận nội dung do AI sinh."
            )
            continue
        normalized: Any
        if section.kind in ("paragraph", "process"):
            normalized = _normalize_paragraph(value)
        elif section.kind == "bullets":
            normalized = _normalize_bullets(value)
        else:
            normalized = _normalize_table(section, value, result.warnings)
        if normalized is None:
            result.warnings.append(
                f"Bỏ nội dung của mục '{key}': sai kiểu dữ liệu cho mục '{section.kind}'."
            )
            continue
        result.content[key] = normalized
    return result


def missing_required_sections(
    template: DocumentTemplate, content: Mapping[str, Any] | None
) -> list[str]:
    """Key các mục bắt buộc chưa có nội dung hợp lệ sau khi chuẩn hóa."""
    cleaned = validate_content(template, content).content
    return [
        section.key
        for section in template.sections
        if section.required and section.accepts_ai_content and not cleaned.get(section.key)
    ]


def _normalize_assumptions(
    assumptions: Sequence[Mapping[str, Any]] | None, warnings: list[str]
) -> list[dict[str, str]]:
    """Giữ các phần tử hồ sơ có status ``assumed``/``skipped`` hợp lệ."""
    items: list[dict[str, str]] = []
    for index, raw in enumerate(assumptions or []):
        if not isinstance(raw, Mapping):
            warnings.append(f"Bỏ giả định thứ {index + 1}: không phải đối tượng.")
            continue
        status = raw.get("status")
        if status not in ASSUMPTION_STATUSES:
            warnings.append(
                f"Bỏ giả định thứ {index + 1}: trạng thái {status!r} không hợp lệ."
            )
            continue
        key = _clean_text(raw.get("key")) or ""
        label = _clean_text(raw.get("label")) or key
        if not label:
            warnings.append(f"Bỏ giả định thứ {index + 1}: thiếu key/label.")
            continue
        value = _cell_text(raw.get("value"))
        items.append({"key": key, "label": label, "status": status, "value": value or ""})
    return items


def _normalize_steps(
    steps: Sequence[Mapping[str, Any]] | None, warnings: list[str]
) -> list[dict[str, str]]:
    rows: list[dict[str, str]] = []
    for index, raw in enumerate(steps or []):
        if not isinstance(raw, Mapping):
            warnings.append(f"Bỏ bước quy trình thứ {index + 1}: không phải đối tượng.")
            continue
        row: dict[str, str] = {}
        for key, _label in PROCESS_STEP_COLUMNS:
            text = _cell_text(raw.get(key))
            row[key] = text if text is not None else ""
        if not row["order"]:
            row["order"] = str(index + 1)
        if row["name"] or row["description"]:
            rows.append(row)
    return rows


# ---------------------------------------------------------------------------
# Các hàm dựng phần tử DOCX
# ---------------------------------------------------------------------------


def _add_toc(document: DocxDocument, language: str) -> None:
    """Chèn field TOC của Word (cấp 1-2) và bật cập nhật field khi mở tệp."""
    document.add_paragraph(_LABELS[language]["toc"], style="TOC Heading")
    paragraph = document.add_paragraph()
    run = paragraph.add_run()

    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = 'TOC \\o "1-2" \\h \\z \\u'
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    placeholder = OxmlElement("w:t")
    placeholder.text = _LABELS[language]["toc_placeholder"]
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for element in (begin, instr, separate, placeholder, end):
        run._r.append(element)

    # Yêu cầu Word cập nhật field (mục lục) khi mở tài liệu.
    settings = document.settings.element
    update = OxmlElement("w:updateFields")
    update.set(qn("w:val"), "true")
    settings.append(update)


def _add_title_page(
    document: DocxDocument,
    *,
    template: DocumentTemplate,
    language: str,
    title: str,
    project_name: str,
    generated_at: datetime,
) -> None:
    text = _LABELS[language]
    heading = document.add_paragraph(title, style="Title")
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle = document.add_paragraph(template.title(language), style="Subtitle")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    date_format = "%Y-%m-%d %H:%M" if language == "en" else "%d/%m/%Y %H:%M"
    for label, value in (
        (text["project"], project_name),
        (text["generated_at"], generated_at.strftime(date_format)),
    ):
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.add_run(f"{label}: ").bold = True
        paragraph.add_run(value)
    document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def _add_note(document: DocxDocument, message: str) -> None:
    """Ghi chú in nghiêng, dùng khi mục chưa có thông tin."""
    run = document.add_paragraph().add_run(message)
    run.italic = True


def _add_table(
    document: DocxDocument, headers: Sequence[str], rows: Sequence[Sequence[str]]
) -> None:
    """Bảng có hàng tiêu đề in đậm, lặp lại khi sang trang."""
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    header_row = table.rows[0]
    # Đánh dấu hàng tiêu đề để Word lặp lại ở đầu mỗi trang.
    tr_pr = header_row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)
    for cell, header in zip(header_row.cells, headers):
        cell.text = ""
        cell.paragraphs[0].add_run(header).bold = True
    for values in rows:
        cells = table.add_row().cells
        for cell, value in zip(cells, values):
            cell.text = value


def _section_number(counters: list[int], level: int) -> str:
    """Tăng bộ đếm theo cấp và trả chuỗi số mục, ví dụ ``2.1``."""
    index = max(level, 1) - 1
    while len(counters) <= index:
        counters.append(0)
    counters[index] += 1
    for deeper in range(index + 1, len(counters)):
        counters[deeper] = 0
    return ".".join(str(value) for value in counters[: index + 1])


def _render_section_body(
    document: DocxDocument,
    section: Section,
    value: Any,
    language: str,
) -> None:
    text = _LABELS[language]
    if not value:
        _add_note(document, text["no_information"])
        return
    if section.kind == "paragraph":
        for paragraph in value:
            document.add_paragraph(paragraph)
    elif section.kind == "bullets":
        for item in value:
            document.add_paragraph(item, style="List Bullet")
    elif section.kind == "table":
        headers = [column.title(language) for column in section.columns]
        rows = [[row[column.key] for column in section.columns] for row in value]
        _add_table(document, headers, rows)


def _render_process(
    document: DocxDocument,
    intro: list[str] | None,
    language: str,
    diagram_png: bytes | None,
    steps: list[dict[str, str]],
    warnings: list[str],
) -> None:
    text = _LABELS[language]
    for paragraph in intro or []:
        document.add_paragraph(paragraph)
    if diagram_png:
        try:
            document.add_picture(io.BytesIO(diagram_png), width=Inches(6))
        # Ảnh hỏng hoặc định dạng lạ: dùng bảng các bước dự phòng thay vì làm hỏng cả tài liệu.
        except Exception as exc:
            warnings.append(f"Không chèn được sơ đồ quy trình ({exc}); dùng bảng các bước.")
        else:
            caption = document.add_paragraph(text["diagram_caption"], style="Caption")
            caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
            return
    if steps:
        headers = [text[label] for _key, label in PROCESS_STEP_COLUMNS]
        rows = [[step[key] for key, _label in PROCESS_STEP_COLUMNS] for step in steps]
        _add_table(document, headers, rows)
        return
    if not intro:
        _add_note(document, text["no_information"])


def _render_assumptions(
    document: DocxDocument, items: list[dict[str, str]], language: str
) -> None:
    text = _LABELS[language]
    if not items:
        _add_note(document, text["no_assumptions"])
        return
    headers = [text["assumption_item"], text["assumption_status"], text["assumption_value"]]
    rows: list[list[str]] = []
    # Giả định trước, vấn đề mở sau; giữ thứ tự gốc trong từng nhóm.
    for status in ("assumed", "skipped"):
        for item in items:
            if item["status"] != status:
                continue
            if status == "assumed":
                value = item["value"] or text["no_information"]
                rows.append([item["label"], text["status_assumed"], value])
            else:
                rows.append([item["label"], text["status_skipped"], text["skipped_value"]])
    _add_table(document, headers, rows)


def render_docx_report(
    template: DocumentTemplate,
    content: Mapping[str, Any] | None,
    assumptions: Sequence[Mapping[str, Any]] | None,
    *,
    language: str = DEFAULT_LANGUAGE,
    title: str,
    project_name: str,
    generated_at: datetime,
    process_diagram_png: bytes | None = None,
    process_steps: Sequence[Mapping[str, Any]] | None = None,
) -> RenderResult:
    """Dựng DOCX và trả kèm danh sách cảnh báo (key lạ, sai kiểu, ảnh hỏng...)."""
    _ensure_language(language)
    validation = validate_content(template, content)
    warnings = list(validation.warnings)
    assumption_items = _normalize_assumptions(assumptions, warnings)
    steps = _normalize_steps(process_steps, warnings)

    document = Document()
    document.core_properties.title = title
    document.core_properties.language = language
    normal_font = document.styles["Normal"].font
    normal_font.name = "Calibri"
    normal_font.size = Pt(11)

    _add_title_page(
        document,
        template=template,
        language=language,
        title=title,
        project_name=project_name,
        generated_at=generated_at,
    )
    _add_toc(document, language)
    document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    counters: list[int] = []
    for section in template.sections:
        number = _section_number(counters, section.level)
        document.add_heading(f"{number}. {section.title(language)}", level=section.level)
        if section.kind == "group":
            continue
        if section.kind == "assumptions":
            _render_assumptions(document, assumption_items, language)
        elif section.kind == "process":
            _render_process(
                document,
                validation.content.get(section.key),
                language,
                process_diagram_png,
                steps,
                warnings,
            )
        else:
            _render_section_body(
                document, section, validation.content.get(section.key), language
            )

    for message in warnings:
        logger.warning("docx_renderer: %s", message)

    buffer = io.BytesIO()
    document.save(buffer)
    return RenderResult(data=buffer.getvalue(), warnings=warnings)


def render_docx(
    template: DocumentTemplate,
    content: Mapping[str, Any] | None,
    assumptions: Sequence[Mapping[str, Any]] | None,
    *,
    language: str = DEFAULT_LANGUAGE,
    title: str,
    project_name: str,
    generated_at: datetime,
    process_diagram_png: bytes | None = None,
    process_steps: Sequence[Mapping[str, Any]] | None = None,
) -> bytes:
    """Dựng DOCX và chỉ trả bytes; cảnh báo được ghi log (xem ``render_docx_report``)."""
    return render_docx_report(
        template,
        content,
        assumptions,
        language=language,
        title=title,
        project_name=project_name,
        generated_at=generated_at,
        process_diagram_png=process_diagram_png,
        process_steps=process_steps,
    ).data
