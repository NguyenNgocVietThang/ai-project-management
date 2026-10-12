"""Renderer XLSX thuần cho báo cáo dự án (task X1, Phase 4).

Hàm `render_project_xlsx` chỉ nhận dữ liệu đã chuẩn bị sẵn (`ProjectReportData`)
và trả về bytes của tệp .xlsx; không truy cập DB, không gọi mạng. Tầng job (X2)
chịu trách nhiệm nạp dữ liệu từ dashboard_service / scheduling_service / CR.

Các nguyên tắc định dạng:
- Mỗi sheet có 2 dòng thông tin dự án, 1 dòng trống, rồi hàng tiêu đề in đậm
  tại `HEADER_ROW`; freeze pane ngay dưới hàng tiêu đề.
- Ngày ghi dưới dạng date/datetime thật (có number_format), số ghi dạng số.
- Task găng được đánh dấu bằng cột chữ "Có/Yes" (màu nền chỉ là phụ trợ).
- Danh sách rỗng vẫn có hàng tiêu đề và một dòng "Không có dữ liệu".
- Chống formula injection: chuỗi bắt đầu bằng = + - @ (và tab/CR) luôn được
  ghi với kiểu dữ liệu chuỗi, không bao giờ thành công thức.
"""
from __future__ import annotations

from collections.abc import Iterable, Sequence
from dataclasses import dataclass, field
from datetime import UTC, date, datetime
from io import BytesIO
from typing import Any, Literal

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

Language = Literal["vi", "en"]

# Hàng tiêu đề cột của mọi bảng chính (1-based). Hàng 1-2 là thông tin dự án.
HEADER_ROW = 4

# Ký tự đầu chuỗi có thể bị bảng tính hiểu là công thức (OWASP CSV injection).
_FORMULA_PREFIXES = ("=", "+", "-", "@", "\t", "\r")

# Bước duyệt mặc định theo step_order của model Approval: 1=BA, 2=PO, 3=PM.
DEFAULT_STEP_ROLES: dict[int, str] = {1: "BA", 2: "PO", 3: "PM"}

_MIN_COL_WIDTH = 8
_MAX_COL_WIDTH = 50

_BOLD = Font(bold=True)
_TITLE_FONT = Font(bold=True, size=13)
_HEADER_FILL = PatternFill("solid", fgColor="DDEBF7")
_CRITICAL_FILL = PatternFill("solid", fgColor="FCE4D6")

_FMT_INT = "0"
_FMT_1 = "0.0"
_FMT_2 = "0.00"
_FMT_MONEY = "#,##0.00"


# ─── Cấu trúc dữ liệu đầu vào ───────────────────────────────────────────────────


@dataclass(frozen=True)
class ProjectInfo:
    """Thông tin dự án in ở đầu mỗi sheet."""

    name: str
    code: str | None = None
    # Ngày/giờ xuất báo cáo; None thì dùng ngày hôm nay.
    exported_at: date | datetime | None = None


@dataclass(frozen=True)
class TaskRow:
    """Một task. Tên field khớp model Task / CPMTask để X2 nạp trực tiếp."""

    id: int
    name: str
    phase: str | None = None
    sprint: str | None = None
    epic: str | None = None
    status: str | None = None
    priority: str | None = None
    assignee: str | None = None
    start_date: date | None = None
    due_date: date | None = None
    progress: float | None = None  # 0-100, giống Task.progress
    estimated_hours: float | None = None
    actual_hours: float | None = None
    duration_days: float | None = None  # CPMTask.duration_days
    early_start: date | None = None
    early_finish: date | None = None
    late_start: date | None = None
    late_finish: date | None = None
    float_days: float | None = None
    is_critical: bool = False


@dataclass(frozen=True)
class CriticalPathData:
    """Đường găng: khớp CPMResponse.critical_path và project_duration_days."""

    task_ids: Sequence[int] = ()
    project_duration_days: float | None = None


@dataclass(frozen=True)
class ReportApprovalStep:
    """Một bước duyệt CR (model Approval)."""

    step_order: int
    decision: str | None = None  # ApprovalStatus: PENDING/APPROVED/REJECTED/...
    role: str | None = None  # BA/PO/PM; None thì suy ra từ step_order
    decided_by: str | None = None
    decided_at: datetime | date | None = None  # Approval.approved_at
    reason: str | None = None  # Approval.comment

    @property
    def resolved_role(self) -> str | None:
        return self.role or DEFAULT_STEP_ROLES.get(self.step_order)


@dataclass(frozen=True)
class ChangeRequestRow:
    """Một CR kèm lịch sử duyệt."""

    id: int
    title: str
    status: str | None = None
    requested_by: str | None = None
    created_at: datetime | date | None = None
    code: str | None = None  # None thì hiển thị "CR-{id}"
    approvals: Sequence[ReportApprovalStep] = ()

    @property
    def display_code(self) -> str:
        return self.code or f"CR-{self.id}"


@dataclass(frozen=True)
class EVAData:
    """Chỉ số giá trị thu được; field nào chưa tính được thì để None."""

    pv: float | None = None
    ev: float | None = None
    ac: float | None = None
    spi: float | None = None
    cpi: float | None = None
    sv: float | None = None
    cv: float | None = None
    bac: float | None = None
    eac: float | None = None


@dataclass(frozen=True)
class BurndownRow:
    """Một điểm burndown. `date` nhận date hoặc chuỗi ISO như BurndownPoint."""

    date: date | str
    ideal: float | None = None
    actual: float | None = None

    @classmethod
    def from_dashboard_point(cls, point: Any) -> BurndownRow:
        """Chuyển BurndownPoint (date: str, remaining, ideal) sang BurndownRow."""
        return cls(date=point.date, ideal=point.ideal, actual=point.remaining)


@dataclass(frozen=True)
class ProjectReportData:
    project: ProjectInfo
    tasks: Sequence[TaskRow] = ()
    critical_path: CriticalPathData = field(default_factory=CriticalPathData)
    change_requests: Sequence[ChangeRequestRow] = ()
    eva: EVAData | None = None
    burndown: Sequence[BurndownRow] = ()


# ─── Nhãn theo ngôn ngữ ─────────────────────────────────────────────────────────

_LABELS: dict[str, dict[str, Any]] = {
    "vi": {
        "sheets": ("Task", "Đường găng", "CR và duyệt", "EVA và Burndown"),
        "project": "Dự án",
        "exported_at": "Ngày xuất",
        "no_data": "Không có dữ liệu",
        "yes": "Có",
        "no": "Không",
        "date_fmt": "dd/mm/yyyy",
        "datetime_fmt": "dd/mm/yyyy hh:mm",
        "task_headers": (
            "ID", "Tên task", "Phase", "Sprint", "Epic", "Trạng thái", "Ưu tiên",
            "Người phụ trách", "Bắt đầu", "Hạn chót", "Tiến độ (%)",
            "Giờ ước tính", "Giờ thực tế", "ES", "EF", "LS", "LF",
            "Float (ngày)", "Găng",
        ),
        "cp_headers": (
            "Thứ tự", "ID", "Tên task", "Thời lượng (ngày)", "ES", "EF", "LS",
            "LF", "Float (ngày)", "Găng",
        ),
        "cp_total": "Tổng thời lượng dự án (ngày)",
        "cr_headers": (
            "Mã CR", "Tiêu đề", "Trạng thái CR", "Người tạo", "Ngày tạo",
            "Bước", "Vai trò", "Quyết định", "Người quyết định", "Thời điểm",
            "Lý do",
        ),
        "eva_headers": ("Chỉ số", "Giá trị"),
        "eva_metrics": (
            ("pv", "PV - Giá trị kế hoạch"),
            ("ev", "EV - Giá trị thu được"),
            ("ac", "AC - Chi phí thực tế"),
            ("spi", "SPI - Chỉ số hiệu quả tiến độ"),
            ("cpi", "CPI - Chỉ số hiệu quả chi phí"),
            ("sv", "SV - Chênh lệch tiến độ"),
            ("cv", "CV - Chênh lệch chi phí"),
            ("bac", "BAC - Ngân sách khi hoàn thành"),
            ("eac", "EAC - Ước tính khi hoàn thành"),
        ),
        "burndown_title": "Burndown",
        "burndown_headers": ("Ngày", "Lý tưởng", "Thực tế (còn lại)"),
    },
    "en": {
        "sheets": ("Task", "Critical path", "Change requests", "EVA & Burndown"),
        "project": "Project",
        "exported_at": "Exported at",
        "no_data": "No data",
        "yes": "Yes",
        "no": "No",
        "date_fmt": "yyyy-mm-dd",
        "datetime_fmt": "yyyy-mm-dd hh:mm",
        "task_headers": (
            "ID", "Task name", "Phase", "Sprint", "Epic", "Status", "Priority",
            "Assignee", "Start", "Due", "Progress (%)",
            "Estimated hours", "Actual hours", "ES", "EF", "LS", "LF",
            "Float (days)", "Critical",
        ),
        "cp_headers": (
            "Order", "ID", "Task name", "Duration (days)", "ES", "EF", "LS",
            "LF", "Float (days)", "Critical",
        ),
        "cp_total": "Project duration (days)",
        "cr_headers": (
            "CR code", "Title", "CR status", "Requested by", "Created at",
            "Step", "Role", "Decision", "Decided by", "Decided at", "Reason",
        ),
        "eva_headers": ("Metric", "Value"),
        "eva_metrics": (
            ("pv", "PV - Planned value"),
            ("ev", "EV - Earned value"),
            ("ac", "AC - Actual cost"),
            ("spi", "SPI - Schedule performance index"),
            ("cpi", "CPI - Cost performance index"),
            ("sv", "SV - Schedule variance"),
            ("cv", "CV - Cost variance"),
            ("bac", "BAC - Budget at completion"),
            ("eac", "EAC - Estimate at completion"),
        ),
        "burndown_title": "Burndown",
        "burndown_headers": ("Date", "Ideal", "Actual (remaining)"),
    },
}

# Định dạng số cho từng chỉ số EVA: chỉ số tỉ lệ 2 chữ số, tiền tệ có phân cách.
_EVA_FORMATS = {"spi": _FMT_2, "cpi": _FMT_2}


def sheet_names(language: Language = "vi") -> tuple[str, str, str, str]:
    """Tên 4 sheet theo ngôn ngữ, theo thứ tự Task, Đường găng, CR, EVA."""
    return _labels(language)["sheets"]


def _labels(language: str) -> dict[str, Any]:
    if language not in _LABELS:
        raise ValueError(f"Ngôn ngữ không hỗ trợ: {language!r}")
    return _LABELS[language]


# ─── Ghi ô an toàn ──────────────────────────────────────────────────────────────


class _SheetWriter:
    """Bọc Worksheet: ghi ô có kiểu, theo dõi độ rộng cột."""

    def __init__(self, ws: Worksheet, labels: dict[str, Any]):
        self.ws = ws
        self.labels = labels
        self._widths: dict[int, int] = {}

    def write(self, row: int, col: int, value: Any, number_format: str | None = None):
        cell = self.ws.cell(row=row, column=col)
        display_len = 0
        if value is None:
            return cell
        if isinstance(value, bool):
            text = self.labels["yes"] if value else self.labels["no"]
            cell.value = text
            display_len = len(text)
        elif isinstance(value, datetime):
            # openpyxl không ghi được datetime có tz: quy về UTC rồi bỏ tzinfo.
            if value.tzinfo is not None:
                value = value.astimezone(UTC).replace(tzinfo=None)
            cell.value = value
            cell.number_format = self.labels["datetime_fmt"]
            display_len = len(self.labels["datetime_fmt"])
        elif isinstance(value, date):
            cell.value = value
            cell.number_format = self.labels["date_fmt"]
            display_len = len(self.labels["date_fmt"])
        elif isinstance(value, int | float):
            cell.value = value
            if number_format:
                cell.number_format = number_format
            display_len = len(f"{value:,.2f}") if isinstance(value, float) else len(str(value))
        else:
            text = ILLEGAL_CHARACTERS_RE.sub("", str(value))
            cell.value = text
            if text.startswith(_FORMULA_PREFIXES):
                # Ép kiểu chuỗi để Excel/LibreOffice không bao giờ coi là công thức.
                cell.data_type = "s"
            display_len = max((len(line) for line in text.splitlines()), default=0)
        self._track(col, display_len)
        return cell

    def header(self, row: int, headers: Iterable[str]) -> None:
        for col, title in enumerate(headers, start=1):
            cell = self.write(row, col, title)
            cell.font = _BOLD
            cell.fill = _HEADER_FILL

    def title_block(self, project: ProjectInfo) -> None:
        """Hai dòng đầu sheet: tên/mã dự án và ngày xuất."""
        name = project.name if not project.code else f"{project.name} ({project.code})"
        label = self.write(1, 1, f"{self.labels['project']}:")
        label.font = _TITLE_FONT
        title = self.write(1, 2, name)
        title.font = _TITLE_FONT
        self.write(2, 1, f"{self.labels['exported_at']}:").font = _BOLD
        self.write(2, 2, project.exported_at or date.today())

    def no_data(self, row: int) -> None:
        self.write(row, 1, self.labels["no_data"])

    def mark_row(self, row: int, n_cols: int, fill: PatternFill) -> None:
        for col in range(1, n_cols + 1):
            self.ws.cell(row=row, column=col).fill = fill

    def freeze_below(self, header_row: int) -> None:
        self.ws.freeze_panes = self.ws.cell(row=header_row + 1, column=1)

    def finalize_widths(self) -> None:
        for col, width in self._widths.items():
            letter = get_column_letter(col)
            self.ws.column_dimensions[letter].width = min(
                max(width + 2, _MIN_COL_WIDTH), _MAX_COL_WIDTH
            )

    def _track(self, col: int, length: int) -> None:
        if length > self._widths.get(col, 0):
            self._widths[col] = length


def _coerce_date(value: date | str | None) -> date | None:
    """Chuyển chuỗi ISO (BurndownPoint.date) thành date; giữ nguyên nếu đã là date."""
    if value is None or isinstance(value, date):
        return value
    try:
        return date.fromisoformat(value[:10])
    except ValueError:
        return None


# ─── Từng sheet ─────────────────────────────────────────────────────────────────


def _write_tasks_sheet(w: _SheetWriter, data: ProjectReportData) -> None:
    headers = w.labels["task_headers"]
    w.title_block(data.project)
    w.header(HEADER_ROW, headers)
    w.freeze_below(HEADER_ROW)
    if not data.tasks:
        w.no_data(HEADER_ROW + 1)
        return
    for offset, task in enumerate(data.tasks, start=1):
        r = HEADER_ROW + offset
        values: list[tuple[Any, str | None]] = [
            (task.id, _FMT_INT),
            (task.name, None),
            (task.phase, None),
            (task.sprint, None),
            (task.epic, None),
            (task.status, None),
            (task.priority, None),
            (task.assignee, None),
            (task.start_date, None),
            (task.due_date, None),
            (task.progress, _FMT_1),
            (task.estimated_hours, _FMT_1),
            (task.actual_hours, _FMT_1),
            (task.early_start, None),
            (task.early_finish, None),
            (task.late_start, None),
            (task.late_finish, None),
            (task.float_days, _FMT_1),
            (bool(task.is_critical), None),
        ]
        for col, (value, fmt) in enumerate(values, start=1):
            w.write(r, col, value, fmt)
        if task.is_critical:
            w.mark_row(r, len(headers), _CRITICAL_FILL)


def _write_critical_path_sheet(w: _SheetWriter, data: ProjectReportData) -> None:
    headers = w.labels["cp_headers"]
    w.title_block(data.project)
    w.header(HEADER_ROW, headers)
    w.freeze_below(HEADER_ROW)
    by_id = {task.id: task for task in data.tasks}
    task_ids = list(data.critical_path.task_ids)
    if not task_ids:
        w.no_data(HEADER_ROW + 1)
        last_row = HEADER_ROW + 1
    else:
        for order, task_id in enumerate(task_ids, start=1):
            r = HEADER_ROW + order
            task = by_id.get(task_id)
            w.write(r, 1, order, _FMT_INT)
            w.write(r, 2, task_id, _FMT_INT)
            if task is not None:
                w.write(r, 3, task.name)
                w.write(r, 4, task.duration_days, _FMT_1)
                w.write(r, 5, task.early_start)
                w.write(r, 6, task.early_finish)
                w.write(r, 7, task.late_start)
                w.write(r, 8, task.late_finish)
                w.write(r, 9, task.float_days, _FMT_1)
            # Mọi task trong danh sách đường găng đều là task găng.
            w.write(r, 10, True)
        last_row = HEADER_ROW + len(task_ids)
    # Dòng tổng thời lượng dự án, cách bảng một dòng trống.
    total_row = last_row + 2
    w.write(total_row, 1, w.labels["cp_total"]).font = _BOLD
    w.write(total_row, 2, data.critical_path.project_duration_days, _FMT_1)


def _write_change_requests_sheet(w: _SheetWriter, data: ProjectReportData) -> None:
    headers = w.labels["cr_headers"]
    w.title_block(data.project)
    w.header(HEADER_ROW, headers)
    w.freeze_below(HEADER_ROW)
    if not data.change_requests:
        w.no_data(HEADER_ROW + 1)
        return
    r = HEADER_ROW
    for cr in data.change_requests:
        steps = sorted(cr.approvals, key=lambda s: s.step_order) or [None]
        # Mỗi bước duyệt một hàng; thông tin CR lặp lại để lọc/sắp xếp được.
        for step in steps:
            r += 1
            w.write(r, 1, cr.display_code)
            w.write(r, 2, cr.title)
            w.write(r, 3, cr.status)
            w.write(r, 4, cr.requested_by)
            w.write(r, 5, cr.created_at)
            if step is not None:
                w.write(r, 6, step.step_order, _FMT_INT)
                w.write(r, 7, step.resolved_role)
                w.write(r, 8, step.decision)
                w.write(r, 9, step.decided_by)
                w.write(r, 10, step.decided_at)
                w.write(r, 11, step.reason)


def _write_eva_sheet(w: _SheetWriter, data: ProjectReportData) -> None:
    w.title_block(data.project)
    w.header(HEADER_ROW, w.labels["eva_headers"])
    w.freeze_below(HEADER_ROW)
    metrics = w.labels["eva_metrics"]
    if data.eva is None:
        w.no_data(HEADER_ROW + 1)
        last_row = HEADER_ROW + 1
    else:
        for offset, (key, label) in enumerate(metrics, start=1):
            r = HEADER_ROW + offset
            w.write(r, 1, label)
            w.write(r, 2, getattr(data.eva, key), _EVA_FORMATS.get(key, _FMT_MONEY))
        last_row = HEADER_ROW + len(metrics)

    # Bảng burndown đặt dưới bảng EVA, cách một dòng trống.
    title_row = last_row + 2
    w.write(title_row, 1, w.labels["burndown_title"]).font = _TITLE_FONT
    bd_header_row = title_row + 1
    w.header(bd_header_row, w.labels["burndown_headers"])
    if not data.burndown:
        w.no_data(bd_header_row + 1)
        return
    for offset, point in enumerate(data.burndown, start=1):
        r = bd_header_row + offset
        w.write(r, 1, _coerce_date(point.date))
        w.write(r, 2, point.ideal, _FMT_1)
        w.write(r, 3, point.actual, _FMT_1)


def burndown_header_row(data: ProjectReportData) -> int:
    """Hàng tiêu đề bảng burndown trong sheet EVA (tiện cho X2 và test)."""
    n_eva_rows = 1 if data.eva is None else len(_LABELS["vi"]["eva_metrics"])
    return HEADER_ROW + n_eva_rows + 3


# ─── API công khai ──────────────────────────────────────────────────────────────


def render_project_xlsx(data: ProjectReportData, *, language: Language = "vi") -> bytes:
    """Sinh tệp XLSX 4 sheet từ dữ liệu dự án; trả về bytes."""
    labels = _labels(language)
    wb = Workbook()
    writers = [
        (_write_tasks_sheet, wb.active),
        (_write_critical_path_sheet, wb.create_sheet()),
        (_write_change_requests_sheet, wb.create_sheet()),
        (_write_eva_sheet, wb.create_sheet()),
    ]
    for (fn, ws), name in zip(writers, labels["sheets"], strict=True):
        ws.title = name
        writer = _SheetWriter(ws, labels)
        fn(writer, data)
        writer.finalize_widths()

    wb.properties.title = data.project.name
    wb.properties.creator = "AI PPM"
    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
